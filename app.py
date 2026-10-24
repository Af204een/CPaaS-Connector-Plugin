import asyncio
import csv
import io
import random
from contextlib import asynccontextmanager
from html import escape
from urllib.parse import urlencode
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from fastapi import FastAPI, Form, Query
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse

from connector import send_message, create_order_message
from database import (
    initialize_database, save_message, get_all_messages, get_message_stats,
    update_message_status, search_messages, save_order, get_all_orders,
    get_message_by_id, get_daily_stats, get_message_history,
    export_messages_csv, export_orders_csv,
)

SIMULATION_DELAY_SECONDS = 12
simulation_tasks = set()
# Build the selector from all regions supported by the phonenumbers library.
# This includes countries and territories; national-format numbers are parsed
# using the selected region, while numbers starting with + are parsed internationally.
import phonenumbers

COUNTRY_NAMES = {
    'AD': 'Andorra',
    'AE': 'United Arab Emirates',
    'AF': 'Afghanistan',
    'AG': 'Antigua and Barbuda',
    'AI': 'Anguilla',
    'AL': 'Albania',
    'AM': 'Armenia',
    'AO': 'Angola',
    'AQ': 'Antarctica',
    'AR': 'Argentina',
    'AS': 'American Samoa',
    'AT': 'Austria',
    'AU': 'Australia',
    'AW': 'Aruba',
    'AX': 'Åland Islands',
    'AZ': 'Azerbaijan',
    'BA': 'Bosnia and Herzegovina',
    'BB': 'Barbados',
    'BD': 'Bangladesh',
    'BE': 'Belgium',
    'BF': 'Burkina Faso',
    'BG': 'Bulgaria',
    'BH': 'Bahrain',
    'BI': 'Burundi',
    'BJ': 'Benin',
    'BL': 'Saint Barthélemy',
    'BM': 'Bermuda',
    'BN': 'Brunei',
    'BO': 'Bolivia',
    'BQ': 'Bonaire, Sint Eustatius and Saba',
    'BR': 'Brazil',
    'BS': 'Bahamas',
    'BT': 'Bhutan',
    'BV': 'Bouvet Island',
    'BW': 'Botswana',
    'BY': 'Belarus',
    'BZ': 'Belize',
    'CA': 'Canada',
    'CC': 'Cocos (Keeling) Islands',
    'CD': 'Democratic Republic of the Congo',
    'CF': 'Central African Republic',
    'CG': 'Republic of the Congo',
    'CH': 'Switzerland',
    'CI': "Côte d'Ivoire",
    'CK': 'Cook Islands',
    'CL': 'Chile',
    'CM': 'Cameroon',
    'CN': 'China',
    'CO': 'Colombia',
    'CR': 'Costa Rica',
    'CU': 'Cuba',
    'CV': 'Cabo Verde',
    'CW': 'Curaçao',
    'CX': 'Christmas Island',
    'CY': 'Cyprus',
    'CZ': 'Czechia',
    'DE': 'Germany',
    'DJ': 'Djibouti',
    'DK': 'Denmark',
    'DM': 'Dominica',
    'DO': 'Dominican Republic',
    'DZ': 'Algeria',
    'EC': 'Ecuador',
    'EE': 'Estonia',
    'EG': 'Egypt',
    'EH': 'Western Sahara',
    'ER': 'Eritrea',
    'ES': 'Spain',
    'ET': 'Ethiopia',
    'FI': 'Finland',
    'FJ': 'Fiji',
    'FK': 'Falkland Islands (Malvinas)',
    'FM': 'Micronesia',
    'FO': 'Faroe Islands',
    'FR': 'France',
    'GA': 'Gabon',
    'GB': 'United Kingdom',
    'GD': 'Grenada',
    'GE': 'Georgia',
    'GF': 'French Guiana',
    'GG': 'Guernsey',
    'GH': 'Ghana',
    'GI': 'Gibraltar',
    'GL': 'Greenland',
    'GM': 'Gambia',
    'GN': 'Guinea',
    'GP': 'Guadeloupe',
    'GQ': 'Equatorial Guinea',
    'GR': 'Greece',
    'GS': 'South Georgia and the South Sandwich Islands',
    'GT': 'Guatemala',
    'GU': 'Guam',
    'GW': 'Guinea-Bissau',
    'GY': 'Guyana',
    'HK': 'Hong Kong',
    'HM': 'Heard Island and McDonald Islands',
    'HN': 'Honduras',
    'HR': 'Croatia',
    'HT': 'Haiti',
    'HU': 'Hungary',
    'ID': 'Indonesia',
    'IE': 'Ireland',
    'IL': 'Israel',
    'IM': 'Isle of Man',
    'IN': 'India',
    'IO': 'British Indian Ocean Territory',
    'IQ': 'Iraq',
    'IR': 'Iran',
    'IS': 'Iceland',
    'IT': 'Italy',
    'JE': 'Jersey',
    'JM': 'Jamaica',
    'JO': 'Jordan',
    'JP': 'Japan',
    'KE': 'Kenya',
    'KG': 'Kyrgyzstan',
    'KH': 'Cambodia',
    'KI': 'Kiribati',
    'KM': 'Comoros',
    'KN': 'Saint Kitts and Nevis',
    'KP': 'North Korea',
    'KR': 'South Korea',
    'KW': 'Kuwait',
    'KY': 'Cayman Islands',
    'KZ': 'Kazakhstan',
    'LA': 'Laos',
    'LB': 'Lebanon',
    'LC': 'Saint Lucia',
    'LI': 'Liechtenstein',
    'LK': 'Sri Lanka',
    'LR': 'Liberia',
    'LS': 'Lesotho',
    'LT': 'Lithuania',
    'LU': 'Luxembourg',
    'LV': 'Latvia',
    'LY': 'Libya',
    'MA': 'Morocco',
    'MC': 'Monaco',
    'MD': 'Moldova',
    'ME': 'Montenegro',
    'MF': 'Saint Martin (French part)',
    'MG': 'Madagascar',
    'MH': 'Marshall Islands',
    'MK': 'North Macedonia',
    'ML': 'Mali',
    'MM': 'Myanmar',
    'MN': 'Mongolia',
    'MO': 'Macao',
    'MP': 'Northern Mariana Islands',
    'MQ': 'Martinique',
    'MR': 'Mauritania',
    'MS': 'Montserrat',
    'MT': 'Malta',
    'MU': 'Mauritius',
    'MV': 'Maldives',
    'MW': 'Malawi',
    'MX': 'Mexico',
    'MY': 'Malaysia',
    'MZ': 'Mozambique',
    'NA': 'Namibia',
    'NC': 'New Caledonia',
    'NE': 'Niger',
    'NF': 'Norfolk Island',
    'NG': 'Nigeria',
    'NI': 'Nicaragua',
    'NL': 'Netherlands',
    'NO': 'Norway',
    'NP': 'Nepal',
    'NR': 'Nauru',
    'NU': 'Niue',
    'NZ': 'New Zealand',
    'OM': 'Oman',
    'PA': 'Panama',
    'PE': 'Peru',
    'PF': 'French Polynesia',
    'PG': 'Papua New Guinea',
    'PH': 'Philippines',
    'PK': 'Pakistan',
    'PL': 'Poland',
    'PM': 'Saint Pierre and Miquelon',
    'PN': 'Pitcairn',
    'PR': 'Puerto Rico',
    'PS': 'Palestine',
    'PT': 'Portugal',
    'PW': 'Palau',
    'PY': 'Paraguay',
    'QA': 'Qatar',
    'RE': 'Réunion',
    'RO': 'Romania',
    'RS': 'Serbia',
    'RU': 'Russia',
    'RW': 'Rwanda',
    'SA': 'Saudi Arabia',
    'SB': 'Solomon Islands',
    'SC': 'Seychelles',
    'SD': 'Sudan',
    'SE': 'Sweden',
    'SG': 'Singapore',
    'SH': 'Saint Helena, Ascension and Tristan da Cunha',
    'SI': 'Slovenia',
    'SJ': 'Svalbard and Jan Mayen',
    'SK': 'Slovakia',
    'SL': 'Sierra Leone',
    'SM': 'San Marino',
    'SN': 'Senegal',
    'SO': 'Somalia',
    'SR': 'Suriname',
    'SS': 'South Sudan',
    'ST': 'Sao Tome and Principe',
    'SV': 'El Salvador',
    'SX': 'Sint Maarten (Dutch part)',
    'SY': 'Syria',
    'SZ': 'Eswatini',
    'TC': 'Turks and Caicos Islands',
    'TD': 'Chad',
    'TF': 'French Southern Territories',
    'TG': 'Togo',
    'TH': 'Thailand',
    'TJ': 'Tajikistan',
    'TK': 'Tokelau',
    'TL': 'Timor-Leste',
    'TM': 'Turkmenistan',
    'TN': 'Tunisia',
    'TO': 'Tonga',
    'TR': 'Türkiye',
    'TT': 'Trinidad and Tobago',
    'TV': 'Tuvalu',
    'TW': 'Taiwan',
    'TZ': 'Tanzania',
    'UA': 'Ukraine',
    'UG': 'Uganda',
    'UM': 'United States Minor Outlying Islands',
    'US': 'United States',
    'UY': 'Uruguay',
    'UZ': 'Uzbekistan',
    'VA': 'Vatican City',
    'VC': 'Saint Vincent and the Grenadines',
    'VE': 'Venezuela',
    'VG': 'Virgin Islands, British',
    'VI': 'Virgin Islands, U.S.',
    'VN': 'Vietnam',
    'VU': 'Vanuatu',
    'WF': 'Wallis and Futuna',
    'WS': 'Samoa',
    'XK': 'Kosovo',
    'YE': 'Yemen',
    'YT': 'Mayotte',
    'ZA': 'South Africa',
    'ZM': 'Zambia',
    'ZW': 'Zimbabwe',
}
REGION_NAMES = {
    'AC': 'Ascension Island',
    'DG': 'Diego Garcia',
    'EA': 'Ceuta and Melilla',
    'EU': 'European Union',
    'IC': 'Canary Islands',
    'TA': 'Tristan da Cunha',
    'UN': 'United Nations',
    'XK': 'Kosovo',
    'ZZ': 'Unknown Region',
}

COUNTRIES = {}

for region in phonenumbers.SUPPORTED_REGIONS:
    calling_code = phonenumbers.country_code_for_region(region)

    if calling_code:
        country_name = COUNTRY_NAMES.get(region, REGION_NAMES.get(region, region))
        COUNTRIES[region] = (
            f"{country_name} ({region}) (+{calling_code})"
        )

COUNTRIES = dict(
    sorted(
        COUNTRIES.items(),
        key=lambda item: (
            item[0] != "IN",
            item[1].lower()
        )
    )
)

initialize_database()




async def simulate_delivery(message_id):
    try:
        await asyncio.sleep(SIMULATION_DELAY_SECONDS)
        current = get_message_by_id(message_id)
        if current and current["status"].lower() == "queued":
            outcome = random.choices(["delivered", "failed"], weights=[95, 5], k=1)[0]
            update_message_status(message_id, outcome)
    except asyncio.CancelledError:
        raise
    except Exception:
        pass


def schedule_simulation(message_id):
    task = asyncio.create_task(simulate_delivery(message_id))
    simulation_tasks.add(task)
    task.add_done_callback(simulation_tasks.discard)


@asynccontextmanager
async def lifespan(app: FastAPI):
    yield
    for task in list(simulation_tasks):
        task.cancel()
    if simulation_tasks:
        await asyncio.gather(*simulation_tasks, return_exceptions=True)


app = FastAPI(title="CPaaS Connector Dashboard", lifespan=lifespan)


def e(value):
    return escape(str(value if value is not None else ""), quote=True)


def build_url(**params):
    clean = {key: value for key, value in params.items() if value not in ("", None, "all")}
    return "/?" + urlencode(clean) if clean else "/"


def status_badge(status):
    status = str(status or "queued").lower()
    return f'<span class="badge {e(status)}">{e(status.title())}</span>'


def svg_trend(daily):
    width, height = 680, 230
    left, right, top, bottom = 42, 16, 18, 38
    plot_w, plot_h = width-left-right, height-top-bottom
    max_value = max([int(row["total"] or 0) for row in daily] + [1])
    points = []
    for index, row in enumerate(daily):
        x = left + (plot_w * index / max(1, len(daily)-1))
        y = top + plot_h - (int(row["total"] or 0) / max_value) * plot_h
        points.append((x, y, row))
    poly = " ".join(f"{x:.1f},{y:.1f}" for x,y,_ in points)
    circles = "".join(
        f'<circle cx="{x:.1f}" cy="{y:.1f}" r="4" fill="#4f46e5"><title>{e(row["day"])}: {row["total"]} messages</title></circle>'
        for x,y,row in points
    )
    labels = "".join(
        f'<text x="{x:.1f}" y="{height-12}" text-anchor="middle" font-size="10" fill="#737b8c">{e(row["day"][-5:])}</text>'
        for x,y,row in points
    )
    grid = "".join(
        f'<line x1="{left}" y1="{top + plot_h*i/3:.1f}" x2="{width-right}" y2="{top + plot_h*i/3:.1f}" stroke="#e8ebf3" stroke-dasharray="4 4"/>'
        for i in range(4)
    )
    return f'<svg viewBox="0 0 {width} {height}" role="img" aria-label="Daily message trend" style="width:100%;height:auto">{grid}<polyline points="{poly}" fill="none" stroke="#4f46e5" stroke-width="3" stroke-linecap="round" stroke-linejoin="round"/>{circles}{labels}</svg>'


@app.get("/", response_class=HTMLResponse)
async def dashboard(
    search: str = "", status: str = "all", notice: str = "",
    order_search: str = "", start_date: str = "", end_date: str = "",
    days: int = Query(7, ge=1, le=30),
):
    stats = get_message_stats()
    messages = search_messages(search, status, start_date, end_date)
    orders = get_all_orders(order_search)
    daily = get_daily_stats(days)
    total, queued = stats["total"], stats["queued"]
    delivered, failed, rate = stats["delivered"], stats["failed"], stats["delivery_rate"]
    resolved = delivered + failed

    def bar_row(label, value, css):
        pct = value / total * 100 if total else 0
        return f'<div class="bar-row"><div class="bar-label"><span>{label}</span><strong>{value}</strong></div><div class="bar-track"><div class="bar-fill {css}" style="width:{pct:.1f}%"></div></div></div>'

    message_rows = ""
    for item in messages:
        message_rows += f"""<tr>
        <td><a class="mono" href="/message/{e(item.get('message_id'))}">{e(item.get('message_id'))}</a></td>
        <td>{e(item.get('phone_number'))}</td><td class="message-cell">{e(item.get('message'))}</td>
        <td>{status_badge(item.get('status'))}</td><td>{e(item.get('created_at'))}</td>
        <td><form class="inline-form" action="/update-status" method="post">
        <input type="hidden" name="message_id" value="{e(item.get('message_id'))}">
        <input type="hidden" name="return_search" value="{e(search)}"><input type="hidden" name="return_status" value="{e(status)}">
        <select name="new_status" aria-label="Update status">
        <option value="queued" {"selected" if item.get("status","").lower()=="queued" else ""}>Queued</option>
        <option value="delivered" {"selected" if item.get("status","").lower()=="delivered" else ""}>Delivered</option>
        <option value="failed" {"selected" if item.get("status","").lower()=="failed" else ""}>Failed</option>
        </select><button class="btn btn-small btn-light">Save</button></form></td></tr>"""
    if not message_rows:
        message_rows = '<tr><td colspan="6" class="empty">No messages match these filters.</td></tr>'

    order_rows = ""
    for order in orders:
        order_rows += f"""<tr><td><a class="mono" href="/order/{e(order.get('order_id'))}">{e(order.get('order_id'))}</a></td>
        <td>₹{float(order.get('amount') or 0):,.2f}</td><td>{e(order.get('phone_number'))}</td>
        <td><a class="mono" href="/message/{e(order.get('message_id'))}">{e(order.get('message_id') or '—')}</a></td>
        <td>{e(order.get('created_at'))}</td></tr>"""
    if not order_rows:
        order_rows = '<tr><td colspan="5" class="empty">No orders found.</td></tr>'

    country_options = "".join(f'<option value="{code}">{label}</option>' for code,label in COUNTRIES.items())
    notice_html = f'<div class="notice">{e(notice)}</div>' if notice else ""
    html = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><meta http-equiv="refresh" content="30">
<title>CPaaS Connector Dashboard</title><style>
:root{{--bg:#f4f7fb;--ink:#172033;--muted:#6b7280;--line:#e6eaf1;--primary:#4f46e5;--shadow:0 10px 30px #1923460f}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font:14px/1.5 Inter,Segoe UI,Arial,sans-serif}}
.topbar{{background:linear-gradient(120deg,#171b3a,#3730a3 70%,#625bf6);color:white;padding:24px clamp(16px,4vw,54px);display:flex;justify-content:space-between;align-items:center;gap:18px;flex-wrap:wrap}}
.brand{{display:flex;gap:13px;align-items:center}}.logo{{width:45px;height:45px;border-radius:14px;background:#ffffff20;border:1px solid #ffffff35;display:grid;place-items:center;font-size:23px;font-weight:800}}
h1{{font-size:23px;margin:0}}.subtitle{{color:#d9dcff;font-size:13px;margin-top:4px}}.live{{border:1px solid #ffffff45;border-radius:999px;padding:8px 12px;font-size:12px}}
.container{{max-width:1500px;margin:auto;padding:26px clamp(14px,3vw,42px) 45px}}.notice{{padding:12px 15px;background:#ecfdf5;color:#166534;border:1px solid #bbf7d0;border-radius:12px;margin-bottom:18px}}
.head{{display:flex;justify-content:space-between;align-items:end;gap:12px;flex-wrap:wrap;margin:25px 0 14px}}.head h2{{font-size:18px;margin:0}}.head p{{margin:4px 0 0;color:var(--muted);font-size:13px}}
.cards{{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:15px}}.card,.panel{{background:white;border:1px solid var(--line);border-radius:17px;box-shadow:var(--shadow);min-width:0}}
.stat{{padding:20px}}.statlabel{{display:flex;justify-content:space-between;color:var(--muted);font-size:11px;font-weight:700;text-transform:uppercase;letter-spacing:.5px}}.ico{{width:36px;height:36px;border-radius:11px;display:grid;place-items:center;background:#eef2ff;font-size:18px}}.num{{font-size:32px;font-weight:750;letter-spacing:-1px;margin-top:12px}}.foot{{font-size:12px;color:var(--muted);margin-top:5px}}
.two{{display:grid;grid-template-columns:1fr 1fr;gap:16px}}.panel{{padding:20px}}.panel h3,.formpanel h3{{font-size:15px;margin:0}}.sub{{font-size:12px;color:var(--muted);margin:4px 0 18px}}
.bar-row{{margin:15px 0}}.bar-label{{display:flex;justify-content:space-between;margin-bottom:6px}}.bar-track{{height:9px;background:#edf0f6;border-radius:20px;overflow:hidden}}.bar-fill{{height:100%;border-radius:20px}}.green{{background:#16a36a}}.amber{{background:#eab308}}.red{{background:#e05260}}
.donutrow{{display:flex;align-items:center;gap:20px;padding:8px 0}}.donut{{width:120px;height:120px;flex-shrink:0;border-radius:50%;display:grid;place-items:center;background:conic-gradient(#16a36a {rate*3.6}deg,#edf0f6 0deg);position:relative}}.donut:before{{content:"";position:absolute;inset:12px;border-radius:50%;background:white}}.donuttext{{position:relative;text-align:center}}.donuttext strong{{display:block;font-size:23px}}.donuttext span{{font-size:10px;color:var(--muted)}}.hint{{font-size:12px;color:var(--muted)}}
.forms{{display:grid;grid-template-columns:1fr 1fr;gap:16px}}.formpanel{{padding:21px}}.formpanel h3{{font-size:16px}}.desc{{color:var(--muted);font-size:12px;margin:4px 0 18px}}
label{{display:block;font-size:12px;font-weight:650;margin:0 0 6px}}input,select,textarea{{width:100%;min-width:0;border:1px solid #dce2ec;border-radius:10px;padding:10px 11px;font:inherit;color:var(--ink);background:white}}textarea{{min-height:88px;resize:vertical}}input:focus,select:focus,textarea:focus{{outline:2px solid #c7d2fe;border-color:#818cf8}}.field{{margin-bottom:13px}}.country-picker{{position:relative}}.country-picker-input{{cursor:text}}.country-picker-list{{display:none;position:absolute;z-index:50;top:calc(100% + 4px);left:0;right:0;max-height:240px;overflow-y:auto;background:#fff;border:1px solid #dce2ec;border-radius:10px;box-shadow:0 12px 28px #19234620}}.country-picker.open .country-picker-list{{display:block}}.country-picker-option{{display:block;width:100%;text-align:left;background:#fff;color:var(--ink);border:0;border-radius:0;padding:9px 11px;font:inherit;cursor:pointer}}.country-picker-option:hover,.country-picker-option:focus{{background:#eef2ff;outline:none}}.country-picker-empty{{padding:10px;color:var(--muted);font-size:12px}}
.btn{{display:inline-flex;justify-content:center;align-items:center;border:0;border-radius:10px;padding:10px 14px;background:var(--primary);color:white;font-weight:700;font-size:12px;text-decoration:none;cursor:pointer;white-space:nowrap}}.btn:hover{{background:#3730a3}}.full{{width:100%}}.light{{background:#f1f3f9;color:#343b50;border:1px solid #e1e5ee}}.small{{padding:6px 8px;font-size:11px}}
.tablepanel{{padding:0;overflow:hidden}}.tablehead{{padding:19px 20px 14px;display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap}}.tablehead h3{{margin:0;font-size:16px}}.tablehead p{{margin:4px 0 0;color:var(--muted);font-size:12px}}.search{{display:flex;align-items:center;gap:7px;flex-wrap:wrap}}.search input{{width:190px}}.search select{{width:125px}}.scroll{{overflow-x:auto}}table{{width:100%;border-collapse:collapse;min-width:900px}}th{{text-align:left;background:#f8f9fc;color:#697386;text-transform:uppercase;font-size:10px;letter-spacing:.5px;padding:12px;border-block:1px solid var(--line);white-space:nowrap}}td{{padding:12px;border-bottom:1px solid #edf0f5;font-size:12px;vertical-align:middle}}tbody tr:hover{{background:#fafbff}}.message-cell{{max-width:260px;overflow-wrap:anywhere;color:#4b5563}}.mono{{font-family:Consolas,monospace;font-size:11px;color:#4f46e5;text-decoration:none}}.badge{{display:inline-block;border-radius:999px;padding:5px 8px;font-size:10px;font-weight:750}}.badge.queued{{background:#fff4d5;color:#956300}}.badge.delivered{{background:#dff8eb;color:#087443}}.badge.failed{{background:#ffe5e7;color:#b42332}}.inline-form{{display:flex;gap:5px;min-width:165px}}.inline-form select{{padding:6px;font-size:11px}}.empty{{text-align:center;padding:27px;color:var(--muted)}}.note{{font-size:11px;color:var(--muted);margin-top:10px}}.footer{{text-align:center;color:#8992a4;font-size:11px;padding-top:25px}}
@media(max-width:900px){{.cards{{grid-template-columns:repeat(2,minmax(0,1fr))}}.two{{grid-template-columns:1fr}}}}@media(max-width:620px){{.cards,.forms{{grid-template-columns:1fr}}h1{{font-size:20px}}.search{{width:100%}}.search input,.search select{{width:100%}}.donut{{width:100px;height:100px}}}}
</style></head><body>
<header class="topbar"><div class="brand"><div class="logo">C</div><div><h1>CPaaS Connector</h1><div class="subtitle">Messaging operations · local simulation environment</div></div></div><div class="live">● &nbsp; LOCAL SIMULATION ACTIVE</div></header>
<main class="container">{notice_html}
<div class="head"><div><h2>Overview</h2><p>Monitor message activity and delivery outcomes.</p></div><div class="hint">Refreshes every 30 seconds</div></div>
<section class="cards">
<article class="card stat"><div class="statlabel">Total messages <span class="ico">✉</span></div><div class="num">{total:,}</div><div class="foot">All saved messages</div></article>
<article class="card stat"><div class="statlabel">Delivered <span class="ico" style="background:#e9f9f1">✓</span></div><div class="num">{delivered:,}</div><div class="foot">Simulated success</div></article>
<article class="card stat"><div class="statlabel">Queued <span class="ico" style="background:#fff6df">◷</span></div><div class="num">{queued:,}</div><div class="foot">Waiting for simulation</div></article>
<article class="card stat"><div class="statlabel">Failed <span class="ico" style="background:#fff0f0">!</span></div><div class="num">{failed:,}</div><div class="foot">Simulated failures</div></article></section>
<div class="head"><div><h2>Analytics</h2><p>Delivery health and message activity over time.</p></div></div>
<section class="two"><article class="panel"><h3>Message status breakdown</h3><p class="sub">Share of all saved messages</p>
{bar_row("Delivered",delivered,"green")}{bar_row("Queued",queued,"amber")}{bar_row("Failed",failed,"red")}
<div class="hint">Failure rate among resolved: {stats['failure_rate']:.1f}%</div></article>
<article class="panel"><h3>Delivery performance</h3><p class="sub">Delivered as a share of resolved outcomes</p><div class="donutrow"><div class="donut"><div class="donuttext"><strong>{rate:.1f}%</strong><span>DELIVERY RATE</span></div></div><div><strong>{delivered:,} delivered</strong><p class="hint">{resolved:,} resolved message(s). Queued messages are excluded until they receive a simulated outcome.</p></div></div></article></section>
<div class="head"><div><h2>Daily message trend</h2><p>Messages created per day. Choose a period to explore activity.</p></div><form class="search" action="/" method="get"><label for="days" style="margin:0">Period</label><select id="days" name="days"><option value="7" {"selected" if days==7 else ""}>Last 7 days</option><option value="14" {"selected" if days==14 else ""}>Last 14 days</option><option value="30" {"selected" if days==30 else ""}>Last 30 days</option></select><button class="btn">Apply</button></form></div>
<section class="panel"><div class="sub">Daily total messages · date labels use MM-DD</div>{svg_trend(daily)}</section>
<div class="head"><div><h2>Create a message</h2><p>Queue a message for free local delivery simulation.</p></div></div>
<section class="forms"><article class="panel formpanel"><h3>Send a message</h3><p class="desc">A new message is simulated after about 12 seconds.</p>
<form action="/send-message" method="post"><div class="field"><label for="country_picker_input">Country / region</label><div class="country-picker" data-select="country_code"><input id="country_picker_input" class="country-picker-input" type="text" role="combobox" aria-expanded="false" aria-autocomplete="list" placeholder="Select or type a country / calling code" autocomplete="off"><div class="country-picker-list" role="listbox"></div><select id="country_code" name="country_code" hidden>{country_options}</select></div></div>
<div class="field"><label for="phone_number">Phone number</label><input id="phone_number" name="phone_number" placeholder="e.g. 9876543210" required></div>
<div class="field"><label for="message">Message text</label><textarea id="message" name="message" maxlength="2000" placeholder="Type your message here..." required></textarea></div><button class="btn full">＋ &nbsp; Queue message</button></form><div class="note">Demo only: no SMS is sent to a real phone.</div></article>
<article class="panel formpanel"><h3>Order notification</h3><p class="desc">Save an order and queue its customer notification.</p><form action="/order-notification" method="post">
<div class="field">
    <label>Order ID</label>
    <input value="Generated automatically" readonly>
    <small>Your Order ID will be generated when you create the order.</small>
</div>
<div class="field"><label for="amount">Order amount (₹)</label><input id="amount" name="amount" type="number" min="0.01" step="0.01" placeholder="e.g. 1499.00" required></div>
<div class="field"><label for="order_country_picker_input">Country / region</label><div class="country-picker" data-select="order_country_code"><input id="order_country_picker_input" class="country-picker-input" type="text" role="combobox" aria-expanded="false" aria-autocomplete="list" placeholder="Select or type a country / calling code" autocomplete="off"><div class="country-picker-list" role="listbox"></div><select id="order_country_code" name="country_code" hidden>{country_options}</select></div></div><div class="field"><label for="order_phone_number">Customer phone number</label><input id="order_phone_number" name="phone_number" placeholder="e.g. 9876543210 or +14155552671" required></div><button class="btn full">＋ &nbsp; Create order notification</button></form><div class="note">Choose the customer country/region for a local-format number, or enter a number in international +country-code format. Demo only: no real SMS is sent.</div></article></section>
<div class="head"><div><h2>Message history</h2><p>Search, filter, export, and update statuses.</p></div><a class="btn" href="/export/messages.xlsx">Export messages Excel</a><a class="btn light" href="/export/messages.csv">Export messages CSV</a></div>
<section class="panel tablepanel"><div class="tablehead"><div><h3>Messages</h3><p>Showing {len(messages)} matching message(s)</p></div>
<form class="search" action="/" method="get"><input name="search" value="{e(search)}" placeholder="Search ID, phone, message">
<select name="status"><option value="all" {"selected" if status.lower()=="all" else ""}>All statuses</option><option value="queued" {"selected" if status.lower()=="queued" else ""}>Queued</option><option value="delivered" {"selected" if status.lower()=="delivered" else ""}>Delivered</option><option value="failed" {"selected" if status.lower()=="failed" else ""}>Failed</option></select>
<input type="date" name="start_date" value="{e(start_date)}" aria-label="Start date"><input type="date" name="end_date" value="{e(end_date)}" aria-label="End date"><button class="btn">Search</button><a class="btn light" href="/">Reset</a></form></div>
<div class="scroll"><table><thead><tr><th>Message ID</th><th>Phone</th><th>Message</th><th>Status</th><th>Created</th><th>Update status</th></tr></thead><tbody>{message_rows}</tbody></table></div></section>
<div class="head"><div><h2>Order history</h2><p>Review saved order notifications.</p></div><a class="btn" href="/export/orders.xlsx">Export orders Excel</a><a class="btn light" href="/export/orders.csv">Export orders CSV</a></div>
<section class="panel tablepanel"><div class="tablehead"><div><h3>Orders</h3><p>Showing {len(orders)} matching order(s)</p></div><form class="search" action="/" method="get"><input name="order_search" value="{e(order_search)}" placeholder="Search order, phone, message ID"><button class="btn">Search orders</button><a class="btn light" href="/">Reset</a></form></div>
<div class="scroll"><table><thead><tr><th>Order ID</th><th>Amount</th><th>Phone</th><th>Message ID</th><th>Created</th></tr></thead><tbody>{order_rows}</tbody></table></div></section>
<div class="footer">CPaaS Connector Prototype · Existing local database preserved · Simulated delivery only</div><script>
(function(){{
  document.querySelectorAll('.country-picker').forEach(function(picker){{
    var select=document.getElementById(picker.dataset.select);
    var input=picker.querySelector('.country-picker-input');
    var list=picker.querySelector('.country-picker-list');
    var options=Array.from(select.options).map(function(option){{return {{value:option.value,text:option.text}};}});
    function render(query){{
      var q=(query||'').trim().toLowerCase();
      var filtered=options.filter(function(item){{
        return item.text.toLowerCase().includes(q)||item.value.toLowerCase().includes(q)||item.text.replace(/\s/g,'').toLowerCase().includes(q.replace(/\s/g,''));
      }});
      list.innerHTML='';
      if(!filtered.length){{var empty=document.createElement('div');empty.className='country-picker-empty';empty.textContent='No matching country or calling code';list.appendChild(empty);return;}}
      filtered.forEach(function(item){{
        var button=document.createElement('button');button.type='button';button.className='country-picker-option';button.setAttribute('role','option');button.textContent=item.text;
        button.addEventListener('click',function(){{select.value=item.value;input.value=item.text;close();}});
        list.appendChild(button);
      }});
    }}
    function open(){{picker.classList.add('open');input.setAttribute('aria-expanded','true');render(input.value);}}
    function close(){{picker.classList.remove('open');input.setAttribute('aria-expanded','false');}}
    input.addEventListener('focus',function(){{input.select();open();}});
    input.addEventListener('click',open);
    input.addEventListener('input',function(){{open();render(input.value);}});
    input.addEventListener('keydown',function(event){{
      if(event.key==='Escape'){{close();input.value=select.options[select.selectedIndex].text;}}
      if(event.key==='Enter'){{var first=list.querySelector('button');if(first){{event.preventDefault();first.click();}}}}
      if(event.key==='ArrowDown'){{var first=list.querySelector('button');if(first){{event.preventDefault();first.focus();}}}}
    }});
    document.addEventListener('click',function(event){{if(!picker.contains(event.target)){{close();if(select.selectedIndex>=0)input.value=select.options[select.selectedIndex].text;}}}});
    if(select.selectedIndex>=0)input.value=select.options[select.selectedIndex].text;
  }});
}})();
</script></main></body></html>"""
    return HTMLResponse(html)


@app.get("/message/{message_id}", response_class=HTMLResponse)
async def message_details(message_id: str):
    item = get_message_history(message_id)
    if not item:
        return HTMLResponse("<h2>Message not found</h2><p><a href='/'>Back to dashboard</a></p>", status_code=404)
    history_rows = "".join(
        f"<tr><td>{e(row['time'])}</td><td>{status_badge(row['status'])}</td><td>{e(row['note'])}</td></tr>"
        for row in item["history"]
    )
    html = f"""<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Message details</title>
<style>body{{font:15px/1.6 Segoe UI,Arial;background:#f4f7fb;color:#172033;margin:0;padding:24px}}main{{max-width:850px;margin:25px auto;background:white;padding:25px;border-radius:16px;box-shadow:0 8px 30px #0001}}h1{{font-size:23px}}dt{{font-weight:700;color:#6b7280;font-size:12px;margin-top:14px}}dd{{margin:3px 0;overflow-wrap:anywhere}}table{{width:100%;border-collapse:collapse;margin-top:18px}}td,th{{text-align:left;padding:11px;border-bottom:1px solid #e6eaf1}}.badge{{padding:4px 8px;border-radius:99px;background:#eee;font-size:12px}}.queued{{background:#fff4d5}}.delivered{{background:#dff8eb}}.failed{{background:#ffe5e7}}a{{color:#4f46e5}}</style></head><body><main>
<a href="/">← Back to dashboard</a><h1>Message details</h1><dl>
<dt>Message ID</dt><dd>{e(item.get('message_id'))}</dd><dt>Phone number</dt><dd>{e(item.get('phone_number'))}</dd>
<dt>Message</dt><dd>{e(item.get('message'))}</dd><dt>Current status</dt><dd>{status_badge(item.get('status'))}</dd>
<dt>Created at</dt><dd>{e(item.get('created_at'))}</dd><dt>Last status update</dt><dd>{e(item.get('updated_at') or item.get('created_at'))}</dd>
<dt>Provider mode</dt><dd>{e(item.get('provider') or 'local-simulation')}</dd></dl>
<h2>Status history (available records)</h2><p>Current database stores the creation time and latest status update, not every intermediate event.</p>
<table><thead><tr><th>Time</th><th>Status</th><th>Event</th></tr></thead><tbody>{history_rows}</tbody></table></main></body></html>"""
    return HTMLResponse(html)



@app.get("/order/{order_id}", response_class=HTMLResponse)
async def order_details(order_id: str):
    matches = get_all_orders(order_id)
    order = next((row for row in matches if row.get("order_id") == order_id), None)
    if not order:
        return HTMLResponse("<h2>Order not found</h2><p><a href='/'>Back to dashboard</a></p>", status_code=404)
    message = get_message_by_id(order.get("message_id")) if order.get("message_id") else None
    message_link = (f'<a href="/message/{e(order.get("message_id"))}">{e(order.get("message_id"))}</a>' if order.get("message_id") else "—")
    message_status = status_badge(message.get("status")) if message else "Not available"
    html = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Order {e(order_id)} details</title>
<style>body{{font:15px/1.6 Segoe UI,Arial,sans-serif;background:#f4f7fb;color:#172033;margin:0;padding:24px}}main{{max-width:850px;margin:25px auto;background:white;padding:28px;border-radius:16px;box-shadow:0 8px 30px #0001}}h1{{font-size:24px}}dt{{font-weight:700;color:#6b7280;font-size:12px;margin-top:16px}}dd{{margin:3px 0;overflow-wrap:anywhere}}a{{color:#4f46e5}}.id{{font-family:monospace;font-size:20px}}</style></head><body><main>
<a href="/">← Back to dashboard</a><h1>Order details</h1><p class="id">{e(order.get('order_id'))}</p><dl>
<dt>Order amount</dt><dd>₹{float(order.get('amount') or 0):,.2f}</dd><dt>Customer phone</dt><dd>{e(order.get('phone_number'))}</dd><dt>Created at</dt><dd>{e(order.get('created_at'))}</dd>
<dt>Associated Message ID</dt><dd>{message_link}</dd><dt>Notification status</dt><dd>{message_status}</dd><dt>Notification text</dt><dd>{e(message.get('message')) if message else 'No associated message found.'}</dd></dl></main></body></html>"""
    return HTMLResponse(html)

def build_excel_file(headers, rows, phone_columns=(), currency_columns=()):
    """Create an Excel workbook with phone numbers stored as text cells."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3730A3")
    for row in rows:
        sheet.append(row)
    for row_index in range(2, sheet.max_row + 1):
        for column_index in phone_columns:
            cell = sheet.cell(row=row_index, column=column_index)
            cell.value = str(cell.value or "")
            cell.number_format = "@"  # Text: keeps + and all digits exactly as saved.
        for column_index in currency_columns:
            sheet.cell(row=row_index, column=column_index).number_format = '"₹"#,##0.00'
    for column_index in range(1, sheet.max_column + 1):
        letter = get_column_letter(column_index)
        max_len = max((len(str(cell.value or "")) for cell in sheet[letter]), default=10)
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 55)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


@app.get("/export/messages.xlsx")
async def export_messages_excel():
    messages = get_all_messages()
    headers = ["Message ID", "Phone Number", "Message", "Status", "Created At", "Updated At", "Provider"]
    rows = [[row.get("message_id", ""), row.get("phone_number", ""), row.get("message", ""),
             row.get("status", ""), row.get("created_at", ""), row.get("updated_at", ""),
             row.get("provider", "local-simulation")] for row in messages]
    content = build_excel_file(headers, rows, phone_columns=(2,))
    return StreamingResponse(iter([content]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="cpaas_messages.xlsx"'})


@app.get("/export/orders.xlsx")
async def export_orders_excel():
    orders = get_all_orders()
    headers = ["Order ID", "Amount", "Phone Number", "Message ID", "Created At"]
    rows = [[row.get("order_id", ""), float(row.get("amount") or 0), row.get("phone_number", ""),
             row.get("message_id", ""), row.get("created_at", "")] for row in orders]
    content = build_excel_file(headers, rows, phone_columns=(3,), currency_columns=(2,))
    return StreamingResponse(iter([content]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="cpaas_orders.xlsx"'})


@app.get("/export/messages.csv")
async def export_messages():
    # UTF-8 BOM helps Excel display Unicode message text correctly (including ₹).
    content = "\ufeff" + export_messages_csv()
    return StreamingResponse(iter([content]), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="cpaas_messages.csv"'})


@app.get("/export/orders.csv")
async def export_orders():
    content = "\ufeff" + export_orders_csv()
    return StreamingResponse(iter([content]), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="cpaas_orders.csv"'})


@app.post("/send-message")
async def send_message_route(phone_number: str = Form(...), message: str = Form(...), country_code: str = Form("IN")):
    result = send_message(phone_number, message, country_code)
    if "message_id" not in result:
        return RedirectResponse(url=build_url(notice=result.get("error", "Unable to create message.")), status_code=303)
    save_message(result)
    schedule_simulation(result["message_id"])
    return RedirectResponse(url=build_url(notice=f"Message {result['message_id']} queued. Local simulation will update it shortly."), status_code=303)


@app.post("/update-status")
async def update_status_route(message_id: str = Form(...), new_status: str = Form(...), return_search: str = Form(""), return_status: str = Form("all")):
    updated = update_message_status(message_id, new_status)
    notice = "Message status updated." if updated else "Unable to update message status."
    return RedirectResponse(url=build_url(search=return_search, status=return_status, notice=notice), status_code=303)



@app.post("/order-notification")
async def order_notification_route(
    amount: float = Form(...),
    phone_number: str = Form(...),
    country_code: str = Form("IN")
):
    if amount <= 0:
        return RedirectResponse(
            url=build_url(
                notice="Order amount must be greater than zero."
            ),
            status_code=303
        )

    # Generate a unique Order ID, such as ORD7854.
    existing_order_ids = {
        order.get("order_id")
        for order in get_all_orders()
    }

    order_id = None

    for _ in range(10000):
        candidate_id = f"ORD{random.randint(1000, 9999)}"

        if candidate_id not in existing_order_ids:
            order_id = candidate_id
            break

    if order_id is None:
        return RedirectResponse(
            url=build_url(
                notice="Unable to generate a unique Order ID. Please try again."
            ),
            status_code=303
        )

    message_text = create_order_message(order_id, amount)

    result = send_message(
        phone_number,
        message_text,
        country_code
    )

    if "message_id" not in result:
        return RedirectResponse(
            url=build_url(
                notice=result.get(
                    "error",
                    "Unable to create order notification."
                )
            ),
            status_code=303
        )

    try:
        save_message(result)

        save_order(
            order_id,
            amount,
            result["phone_number"],
            result["message_id"],
            result["created_at"]
        )

    except Exception:
        return RedirectResponse(
            url=build_url(
                notice="Could not save the order. Please try again."
            ),
            status_code=303
        )

    schedule_simulation(result["message_id"])

    return RedirectResponse(
        url=build_url(
            notice=(
                f"Order {order_id} created successfully. "
                "Notification queued for local simulation."
            )
        ),
        status_code=303
    )